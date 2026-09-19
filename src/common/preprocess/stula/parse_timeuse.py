"""
社会生活基本調査 (令和3年, 調査票A) 公表「生活時間編」Excel表をCSVに変換するパーサ

設計: src/models/DDPM_Aggregate_Simple/docs/Stage2_design.md §9.7（実装項目 17）

**何のための表か。**時間帯編（`parse_timeband.py`）が与える教師 `A*` は時刻別行動者率で、
Stage 2 の損失そのものである。したがって `A*` への適合は循環しており、それだけでは主張に
ならない（§9.2）。この表が与える**日次行動者率**は教師が縛らない量で、**日本の公表値に
照らせる唯一の非循環な軸**である。

データ入手 (e-Stat 統計コード 00200533 / 生活時間編):
    data/raw/opened/STULA2021/<statInfId>.xlsx に配置する。
        000032224198: 第70-3表 曜日,男女,ふだんの就業状態,ふだんの健康状態,年齢,
                      行動の種類別行動者率(15歳以上) - 全国,都道府県
    curl はブラウザ相当 UA + Referer で通る（一覧 lid=000001298057）。

★時間帯編と**向きが逆**である。ここが流用できない理由:

```mermaid
flowchart LR
    subgraph TB["時間帯編 第8-1表 (parse_timeband.py)"]
        TB1["行 = 曜日×地域×男女×就業×<b>行動</b>×年齢"]
        TB2["列 = 時刻区分 96 (s0..s95)"]
        TB1 --> TB3["1 行 = 1 行動 × 96 時刻"]
        TB2 --> TB3
    end
    subgraph TU["生活時間編 第70-3表 (このファイル)"]
        TU1["行 = 曜日×地域×男女×就業×健康×年齢<br/>＝<b>層</b>そのもの"]
        TU2["列 = <b>行動</b> 20 (a01..a20)"]
        TU1 --> TU3["1 行 = 1 層 × 20 行動"]
        TU2 --> TU3
    end
```

**★行が層そのものなので `サンプルサイズ` は普通の列でよい。**時間帯編では行動が行キー
だったため、層の属性（推定人口・サンプルサイズ）が `行動=00_総数` 行にしか入らず、
層粒度の別表 `<name>_layers.csv` へ逃がす必要があった（§3.2）。こちらは 1 行 1 層なので
その回避が要らない。

**★12 分類への集約はここでやらない。**日次行動者率は**分類の和で足せない**:

    P(∃s: x_s ∈ C) ≠ Σ_{c∈C} P(∃s: x_s = c)

`crosswalk_atus_stula.stula_to_common` は単純和なので、時刻別行動者率には正しいが
日次行動者率には**誤り**である。common12 と 1 対 1 なのは 7 活動だけで、残る 5 活動は
上下限でしか言えない（§9.7）。その処理は `common12_bounds` が持つ。

処理:
    1. 行9からキー列、行7から行動20列、行5から推定人口列とサンプルサイズ列を自動検出
    2. 欠損マーカーの区別 (e-Stat 慣行、parse_timeband.py と同じ):
        '-'            = 該当数字なし = 行動者ゼロ -> 0.0
        '…' / '...' / 'X' = 非公表        -> NaN
    3. 率は % のまま保持 (0..100)
    4. 再掲列 (R1-R3 の1次/2次/3次活動) は落とす。20 分類の再集計であって独立な情報ではない

出力: data/processed/stula/<name>.csv
    キー6列 + population_k + sample_size + a01..a20
"""

from pathlib import Path
from typing import cast

import numpy as np
import pandas as pd
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[4]
RAW_DIR = REPO_ROOT / "data" / "raw" / "opened" / "STULA2021"
OUT_DIR = REPO_ROOT / "data" / "processed" / "stula"

# 処理対象: statInfId -> 出力名
TABLES = {
    "000032224198": "timeuse_participation",       # 第70-3表 行動者率
}

N_ACT = 20            # 社基調の行動20分類。再掲 R1-R3 は含めない
BLOCK_HEAD_ROW = 5    # ブロック見出し (推定人口 / 行動者率 / サンプルサイズ)
ACT_LABEL_ROW = 7     # 行動ラベルの行 ('01_睡眠' 等)
HEADER_KEYS_ROW = 9   # キー列名の行
DATA_START_ROW = 10
POP_HEAD = "推定人口"
SAMPLE_SIZE_HEAD = "サンプルサイズ"

# キー列名の正規化 (表側の日本語ヘッダ -> 出力列名)
KEY_RENAME = {
    "曜日": "daytype",
    "地域区分": "region",
    "男女": "gender",
    "ふだんの就業状態": "employment",
    "ふだんの健康状態": "health",
    "年齢": "age_class",
}


def _to_num(v) -> float:
    """数値へ。非数値は一律 NaN（推定人口・サンプルサイズ列用）。"""
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def _to_rate(v) -> float:
    """率セル。'-' は「該当数字なし」= 行動者ゼロ -> 0.0、非公表のみ NaN。

    Note:
        ★この区別が最も効く（§3.3 P1）。非公表を 0 にすると教師に偽の観測を足す
          ことになる。実測では 28 群（平日・全国・健康総数）の 560 セルに
          どちらのマーカーも現れないが、地域や細分区分へ広げると両方出る。
    """
    if isinstance(v, str) and v.strip() == "-":
        return 0.0
    return _to_num(v)


def parse_table(path: Path) -> pd.DataFrame:
    """生活時間編の行動者率表を 1 行 1 層の wide 形式へ変換する。

    Args:
        path: 生活時間編の xlsx（第70-3表）

    Returns:
        キー6列 + population_k + sample_size + a01..a20 の DataFrame。
        率は % のまま (0..100)、非公表は NaN

    Raises:
        AssertionError: 行動列が 20 でない場合、推定人口列またはサンプルサイズ列を
            見つけられない場合
    """
    wb = load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    head = [next(rows) for _ in range(DATA_START_ROW - 1)]

    key_cols = {c: str(v).strip() for c, v in enumerate(head[HEADER_KEYS_ROW - 1])
                if v is not None and str(v).strip() != ""}
    # 行動列: 行7 が 'NN_ラベル' で NN が 01..20 のもの。再掲 R1-R3 はここで落ちる
    act_cols = {}
    for c, v in enumerate(head[ACT_LABEL_ROW - 1]):
        if v is None:
            continue
        code = str(v).split("_")[0]
        if code.isdigit() and 1 <= int(code) <= N_ACT:
            act_cols[c] = int(code)
    assert len(act_cols) == N_ACT, f"行動列が20でない: {len(act_cols)} ({path.name})"

    def _head_col(label: str) -> int:
        hit = [c for c, v in enumerate(head[BLOCK_HEAD_ROW - 1])
               if v is not None and str(v).strip() == label]
        assert len(hit) == 1, f"'{label}' 列が1つでない: {hit} ({path.name})"
        return hit[0]

    pop_col, size_col = _head_col(POP_HEAD), _head_col(SAMPLE_SIZE_HEAD)

    records = []
    for row in rows:
        if row[list(key_cols)[0]] is None:
            continue
        rec: dict[str, str | float] = {
            KEY_RENAME.get(name, name): str(row[c]).strip()
            for c, name in key_cols.items() if name.strip() != ""}
        rec["population_k"] = _to_num(row[pop_col])
        rec["sample_size"] = _to_num(row[size_col])
        for c, code in act_cols.items():
            rec[f"a{code:02d}"] = _to_rate(row[c])
        records.append(rec)
    return fix_no_sample_layers(pd.DataFrame(records))


def fix_no_sample_layers(df: pd.DataFrame) -> pd.DataFrame:
    """標本が 1 人も居ない層の率を NaN へ戻す。

    Note:
        ★`'-' -> 0.0` の例外処理である。標本なしの層は 20 行動すべてが '-' で埋まり、
          素直に読むと「その層の人は 1 日に一度も眠らない」という偽の観測になる。
        ★**判別は `サンプルサイズ` 列で厳密にできる。**時間帯編の同じ処理
          (`parse_timeband.fix_rate_semantics` (b)) は層のスロット合計が 0 かどうかで
          推定していたが、この表はサンプルサイズを持つので推定が要らない。
          実測でも両者は完全に一致した（NaN の 49,754 行が、ちょうど 20 行動すべて 0 の
          行と一対一）。
        ★推定人口列は触らない。標本が無くても推計人口は公表されうるため。

    Args:
        df: 読み取り直後の DataFrame

    Returns:
        標本なし層の a01..a20 を NaN にした DataFrame
    """
    act_cols = [f"a{i:02d}" for i in range(1, N_ACT + 1)]
    empty = df["sample_size"].isna() | (df["sample_size"] <= 0)
    if empty.any():
        df.loc[empty, act_cols] = np.nan
        print(f"  標本なし層 (サンプルサイズ無し): {int(empty.sum())} 層を NaN 化")
    return df


def validate(df: pd.DataFrame, name: str) -> None:
    """公表値として素直に読めるかを確かめる。

    Note:
        ★時間帯編の `validate` とは別の検査である。あちらは「主行動は排他的なので
          20 分類のスロット合計が 100%」を見るが、**日次行動者率は排他的ではない**
          （1 日に睡眠も食事も仕事もする）。合計は 100% を大きく超えるのが正常で、
          同じ検査を当てると必ず落ちる。ここで見るのは率の定義域と、睡眠がほぼ全員に
          立つこと（1 日に 1 度も眠らない人はほぼ居ない）である。

    Args:
        df: parse_table の戻り値
        name: 表の名前（メッセージ用）
    """
    act_cols = [f"a{i:02d}" for i in range(1, N_ACT + 1)]
    vals = df[act_cols].to_numpy(dtype=float)
    finite = vals[~np.isnan(vals)]
    assert finite.min() >= 0.0 and finite.max() <= 100.0, \
        f"率が [0,100] の外にある: {finite.min()}〜{finite.max()}"
    total = np.nansum(vals, axis=1)
    sleep = df["a01"].to_numpy(dtype=float)
    ok = ~np.isnan(sleep)
    print(f"  [{name}] 率の範囲 {finite.min():.1f}〜{finite.max():.1f}%  "
          f"20分類の合計 {total.min():.0f}〜{total.max():.0f}%（排他でないので100%超が正常）")
    print(f"  [{name}] 公表行 {int(ok.sum())} / {len(df)}")

    # ★読み取り位置の検査は、パイプラインが実際に使う行で行う。
    #   全体へ一律に掛けてはいけない: 県別 × 健康「良くない」には n=10〜53 の実在する
    #   小標本があり、そこでは睡眠が 64〜89% まで落ちる。これは公表値であって
    #   読み取り誤りではないので、閾値を全体に当てると正しい表を弾いてしまう。
    used = cast(pd.DataFrame, df[(df["region"] == "00_全国") & (df["health"] == "0_総数")])
    s_used = used[["a01"]].to_numpy(dtype=float).ravel()
    s_ok = s_used[~np.isnan(s_used)]
    print(f"  [{name}] 全国・健康総数の睡眠 a01: {s_ok.min():.1f}〜{s_ok.max():.1f}% "
          f"({len(s_ok)} 行)")
    assert s_ok.min() > 95.0, \
        "全国・健康総数で睡眠の日次行動者率が95%を割る（列の読み取り位置を疑う）"
    thin = int(((sleep < 90.0) & ok).sum())
    print(f"  [{name}] 参考: 睡眠<90% の行 {thin} 件（いずれも小標本の細分区分）")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for stat_id, name in TABLES.items():
        path = RAW_DIR / f"{stat_id}.xlsx"
        if not path.exists():
            print(f"[skip] {path} が無い (e-Stat statInfId={stat_id} をDLして配置)")
            continue
        print(f"parsing {path.name} ({name}) ...")
        df = parse_table(path)
        key_desc = {c: df[c].nunique() for c in df.columns
                    if not c.startswith("a") and c not in ("population_k", "sample_size")}
        print(f"  rows={len(df)}  keys={key_desc}")
        validate(df, name)
        out = OUT_DIR / f"{name}.csv"
        df.to_csv(out, index=False)
        print(f"  -> {out}")


if __name__ == "__main__":
    main()
