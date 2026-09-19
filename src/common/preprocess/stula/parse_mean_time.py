"""
社会生活基本調査 (令和3年, 調査票A) 公表「平均時刻編」Excel表をCSVに変換するパーサ

設計: src/models/DDPM_Aggregate_Simple/docs/Stage2_design.md §9.7（実装項目 17）

**何のための表か。**`src/eval/stula_derived_times.py` が生成物へ当てる起床・就寝の規則は、
公表統計が実際に使っている操作的定義である。その規則で日本の実測値がどこに出ているかを
与えるのがこの表で、**派生時刻の突合先**になる。教師（時刻別行動者率）とは無関係な量なので
非循環である（§9.2）。

データ入手 (e-Stat 統計コード 00200533 / 平均時刻編、一覧 lid=000001298059):
    data/raw/opened/STULA2021/<statInfId>.xlsx に配置する。
        000032224374: 第3-2表  曜日,起床,男女,ふだんの就業状態,ライフステージ,年齢,
                      時刻区分別行動者数（構成比）(15歳以上) - 全国
        000032224412: 第24-1表 曜日,就寝,男女,ふだんの就業状態,ライフステージ,年齢,
                      時刻区分別行動者数（構成比）(15歳以上) - 全国
    curl はブラウザ相当 UA + Referer で通る。

★**起床の表選びに落とし穴がある。**「起床」の表は 第1-3表〜第7表 と多数あるが、
  28 群に必要な `ふだんの就業状態 × 年齢` の両方を持つのは **第3-2表だけ**である。
  よく見つかる 第3-1表 は就業状態を持たず（かつ 10 歳以上）、第1-4表 は就業状態を持つが
  年齢の代わりにスマートフォン使用時間を持つ。表題だけで選ぶと軸が足りない。

★**`平均時刻` は 0〜36 時の軸で公表されている。**就寝の実測値は `21:24`〜`24:43` の範囲で、
  深夜 0:12 の就寝は **`24:12`** と書かれる（`0:12` ではない）。24 時間で折り返さないこと。
  `src/eval/stula_derived_times.bed_time` が同じ軸を返すのはこの表に合わせるためである。

★**`行動者率` は「時刻が定まった人の割合」である。**起床・就寝が「不詳」になる人がいる
  （規則を満たす睡眠が無い、変則勤務など）。実測は 99% 前後で、生成側の不詳率と
  比べる対象になる。単なる付随情報ではない。

処理:
    1. 行9からキー列、行5から推定人口・行動者率・平均時刻の列を自動検出
    2. `平均時刻` の "H:MM" を時（float）へ。非公表（'…' 等）は NaN
    3. 時刻区分別の構成比（行動者数の分布）は**読まない**。突合に使うのは平均時刻と
       行動者率で、分布は表ごとに階級の切り方が違う（起床42階級・就寝34階級、両端が
       開区間）ため別の設計が要る

出力: data/processed/stula/<name>.csv
    キー6列 + population_k + actor_rate + mean_time_hour
"""

from pathlib import Path
from typing import cast

import numpy as np
import pandas as pd
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[4]
RAW_DIR = REPO_ROOT / "data" / "raw" / "opened" / "STULA2021"
OUT_DIR = REPO_ROOT / "data" / "processed" / "stula"

# 処理対象: statInfId -> (出力名, 期待される行動の名前)
TABLES = {
    "000032224374": ("meantime_wake", "起床"),
    "000032224412": ("meantime_bed", "就寝"),
}

BLOCK_HEAD_ROW = 5     # ブロック見出し (推定人口 / 行動者率 / 平均時刻 / 行動者数（構成比）)
HEADER_KEYS_ROW = 9
DATA_START_ROW = 10
POP_HEAD = "推定人口"
RATE_HEAD = "行動者率"
MEAN_HEAD = "平均時刻"

KEY_RENAME = {
    "曜日": "daytype",
    "地域区分": "region",
    "男女": "gender",
    "ふだんの就業状態": "employment",
    "ライフステージ": "life_stage",
    "年齢": "age_class",
}


def _to_num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


def parse_mean_time_hour(v) -> float:
    """公表の "H:MM" を時（float）へ。非公表は NaN。

    Note:
        ★24 時間で折り返さない。就寝は 0〜36 時の軸で公表されており、`24:12` は
          25 時間目の 12 分、すなわち 24.2 である。`% 24` を掛けると深夜の就寝が
          0.2 になり、平均が壊れる。
        ★openpyxl が時刻セルを `datetime.time` として返す場合も拾う。ただし
          `time` は 24 時を表現できないので、その経路に落ちるのは起床側だけである。

    Args:
        v: セルの値。"6:43" / "24:12" のような文字列、または datetime.time

    Returns:
        0 時からの時間, dtype=float。非公表・解釈不能は NaN
    """
    if v is None:
        return np.nan
    if hasattr(v, "hour") and hasattr(v, "minute"):
        return float(v.hour) + float(v.minute) / 60.0
    s = str(v).strip()
    if ":" not in s:
        return np.nan
    hh, _, mm = s.partition(":")
    try:
        return float(int(hh)) + float(int(mm)) / 60.0
    except ValueError:
        return np.nan


def parse_table(path: Path) -> pd.DataFrame:
    """平均時刻編の表を 1 行 1 層へ変換する。

    Args:
        path: 平均時刻編の xlsx

    Returns:
        キー6列 + population_k + actor_rate + mean_time_hour の DataFrame

    Raises:
        AssertionError: 推定人口・行動者率・平均時刻の列が 1 つに定まらない場合
    """
    wb = load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    head = [next(rows) for _ in range(DATA_START_ROW - 1)]

    key_cols = {c: str(v).strip() for c, v in enumerate(head[HEADER_KEYS_ROW - 1])
                if v is not None and str(v).strip() != ""}

    def _head_col(label: str) -> int:
        hit = [c for c, v in enumerate(head[BLOCK_HEAD_ROW - 1])
               if v is not None and str(v).strip() == label]
        assert len(hit) == 1, f"'{label}' 列が1つでない: {hit} ({path.name})"
        return hit[0]

    pop_col, rate_col, mean_col = (_head_col(POP_HEAD), _head_col(RATE_HEAD),
                                   _head_col(MEAN_HEAD))

    records = []
    for row in rows:
        if row[list(key_cols)[0]] is None:
            continue
        rec: dict[str, str | float] = {KEY_RENAME.get(name, name): str(row[c]).strip()
                                       for c, name in key_cols.items()}
        rec["population_k"] = _to_num(row[pop_col])
        rec["actor_rate"] = _to_num(row[rate_col])
        rec["mean_time_hour"] = parse_mean_time_hour(row[mean_col])
        records.append(rec)
    return pd.DataFrame(records)


def validate(df: pd.DataFrame, name: str, act: str) -> None:
    """読み取り位置と軸の検査。

    Note:
        ★28 群に必要な軸が揃っているかをここで落とす。起床の表は候補が多く、
          就業状態や年齢を持たないものを取り違えやすい（モジュール docstring 参照）。
        ★時刻の範囲も見る。起床は 0〜12 時、就寝は 17〜36 時に入るはずで、
          外に出たら列の取り違えか折り返しの誤りである。

    Args:
        df: parse_table の戻り値
        name: 出力名（メッセージ用）
        act: 期待される行動（"起床" / "就寝"）
    """
    for col in ("employment", "age_class", "gender", "daytype"):
        assert col in df.columns, f"{name}: 軸 {col} が無い。表を取り違えている"
    need = {"1_有業者", "2_無業者"}
    assert need <= set(df["employment"]), \
        f"{name}: 就業状態に {need - set(df['employment'])} が無い"
    mt = df["mean_time_hour"].to_numpy(dtype=float)
    ok = ~np.isnan(mt)
    lo, hi = (0.0, 12.0) if act == "起床" else (17.0, 36.0)
    bad = ok & ((mt < lo) | (mt >= hi))
    print(f"  [{name}] {act}: 平均時刻 {mt[ok].min():.2f}〜{mt[ok].max():.2f} 時 "
          f"(公表 {int(ok.sum())} / {len(df)} 行)")
    assert not bad.any(), \
        (f"{name}: 平均時刻が {lo}〜{hi} 時の外に {int(bad.sum())} 行ある "
         f"（列の取り違えか 24 時での折り返し）")
    rate = df["actor_rate"].to_numpy(dtype=float)
    r_ok = ~np.isnan(rate)
    print(f"  [{name}] 行動者率（時刻が定まった割合）: "
          f"{rate[r_ok].min():.1f}〜{rate[r_ok].max():.1f}%")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for stat_id, (name, act) in TABLES.items():
        path = RAW_DIR / f"{stat_id}.xlsx"
        if not path.exists():
            print(f"[skip] {path} が無い (e-Stat statInfId={stat_id} をDLして配置)")
            continue
        print(f"parsing {path.name} ({name}) ...")
        df = parse_table(path)
        keys = {c: cast(pd.Series, df[c]).nunique() for c in df.columns
                if c in KEY_RENAME.values()}
        print(f"  rows={len(df)}  keys={keys}")
        validate(df, name, act)
        out = OUT_DIR / f"{name}.csv"
        df.to_csv(out, index=False)
        print(f"  -> {out}")


if __name__ == "__main__":
    main()
